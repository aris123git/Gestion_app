"""Tests : exports de rapports complets (Excel / PDF)."""

from __future__ import annotations

import os
import shutil
import tempfile
import unittest
from datetime import date, timedelta

_DATA_DIR = tempfile.mkdtemp(prefix="gestion_report_export_")
os.environ["GESTION_DATA_DIR"] = _DATA_DIR
os.environ["NEXAPOS_SKIP_ACTIVATION"] = "1"

from app import config  # noqa: E402
from app.controllers.product_controller import ProductController  # noqa: E402
from app.controllers.report_controller import ReportController  # noqa: E402
from app.controllers.sale_controller import CartLine, PaymentLine, SaleController  # noqa: E402
from app.database.connection import engine, init_database  # noqa: E402
from app.database.seed import seed_all  # noqa: E402
from app.reports.excel_report import export_report_excel  # noqa: E402
from app.reports.pdf_report import export_report_pdf  # noqa: E402


class ReportExportTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        init_database()
        seed_all()
        cls.product = ProductController.create(
            {"name": "Riz", "sale_price": 1000, "purchase_price": 600, "quantity": 20, "tax_rate": 18}
        )
        SaleController.create_sale(
            [CartLine(cls.product.id, "Riz", 1000, 2)], [PaymentLine("Espèces", 2000)], amount_received=2000
        )
        cls.start, cls.end = date.today() - timedelta(days=3), date.today()
        cls.report = ReportController.build(cls.start, cls.end)

    @classmethod
    def tearDownClass(cls) -> None:
        engine.dispose()
        shutil.rmtree(config.DATA_DIR, ignore_errors=True)

    def test_excel_has_detail_sheets(self) -> None:
        import openpyxl

        path = export_report_excel(self.report, path=config.EXPORT_DIR / "t.xlsx")
        wb = openpyxl.load_workbook(path)
        for name in ("Synthèse", "Lignes vendues", "Ventes par jour", "Par caissier", "TVA par taux"):
            self.assertIn(name, wb.sheetnames)
        self.assertEqual(wb["Lignes vendues"]["C2"].value, "Riz")

    def test_excel_without_profit_hides_margins(self) -> None:
        import openpyxl

        path = export_report_excel(self.report, path=config.EXPORT_DIR / "np.xlsx", include_profits=False)
        names = openpyxl.load_workbook(path).sheetnames
        self.assertNotIn("Marges produits", names)
        self.assertNotIn("Dépenses", names)

    def test_sections_filter(self) -> None:
        import openpyxl

        path = export_report_excel(self.report, path=config.EXPORT_DIR / "s.xlsx", sections={"daily"})
        names = openpyxl.load_workbook(path).sheetnames
        self.assertIn("Ventes par jour", names)
        self.assertNotIn("Lignes vendues", names)

    def test_pdf_generated(self) -> None:
        path = export_report_pdf(self.report, path=config.EXPORT_DIR / "t.pdf")
        self.assertGreater(path.stat().st_size, 1000)


if __name__ == "__main__":
    unittest.main()
