"""Exports Excel / PDF de la comptabilité (partie double)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from app import config
from app.reports.excel_report import _sheet
from app.reports.pdf_report import _footer, _table
from app.services import accounting_service as acc
from app.services import settings_service
from app.services.chart_of_accounts import account_label
from app.utils.helpers import format_money

BALANCE_HEADERS = ["Compte", "Intitulé", "Ouverture débit", "Ouverture crédit", "Mouvements débit", "Mouvements crédit", "Solde débit", "Solde crédit"]


def _sheet_rows(data: dict) -> list[tuple]:
    sheet = data["sheet"]
    rows = [("ACTIF", "", None)] + [(a, l, v) for a, l, v in sheet["assets"]] + [("", "TOTAL ACTIF", sheet["total_assets"]), ("PASSIF", "", None)]
    rows += [(a, l, v) for a, l, v in sheet["liabilities"]] + [("", "TOTAL PASSIF", sheet["total_liabilities"])]
    return rows


def export_accounting_excel(data: dict, path: str | Path | None = None) -> Path:
    config.ensure_directories()
    shop = settings_service.get_shop_info()
    if path is None:
        path = config.EXPORT_DIR / f"comptabilite_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    path = Path(path)
    wb = Workbook()
    rows = acc.income_rows(data)
    sh = _sheet(wb, "Compte de résultat", ["Libellé", f"Montant ({shop.currency or 'FCFA'})"], [(l, a) for l, a, _b in rows], money_cols=(2,), first=True)
    for r, (_l, _a, bold) in enumerate(rows, start=2):
        if bold:
            sh.cell(row=r, column=1).font = Font(bold=True)
            sh.cell(row=r, column=2).font = Font(bold=True)
    sh.column_dimensions["A"].width = 55
    sh.column_dimensions["B"].width = 22
    sh.append([])
    sh.append([f"Période : {data['start']:%d/%m/%Y} au {data['end']:%d/%m/%Y}"])
    bs = _sheet(wb, "Bilan", ["Compte", "Intitulé", "Montant"], _sheet_rows(data), money_cols=(3,))
    bs.column_dimensions["B"].width = 48
    for row in bs.iter_rows(min_row=2):
        if row[0].value in ("ACTIF", "PASSIF") or str(row[1].value).startswith("TOTAL"):
            for c in row:
                c.font = Font(bold=True)
    _sheet(wb, "Journal", acc.JOURNAL_HEADERS, data["journal"], money_cols=(7, 8), total_cols=(7, 8))
    gl = wb.create_sheet("Grand livre")
    gl.append(["Compte", "Intitulé", "Date", "Journal", "Pièce", "Libellé", "Débit", "Crédit", "Solde"])
    for acct, lines in data["ledger"].items():
        for date_, jc, piece, label, debit, credit, run in lines:
            gl.append([acct, account_label(acct), date_, jc, piece, label, debit, credit, run])
    from app.reports.excel_report import _autosize, _style_header

    _style_header(gl, 9)
    gl.freeze_panes = "A2"
    gl.auto_filter.ref = f"A1:I{gl.max_row}"
    for r in range(2, gl.max_row + 1):
        for c in (7, 8, 9):
            gl.cell(row=r, column=c).number_format = "#,##0.###"
    _autosize(gl)
    _sheet(wb, "Balance", BALANCE_HEADERS, data["balance"], money_cols=(3, 4, 5, 6, 7, 8), total_cols=(3, 4, 5, 6, 7, 8))
    vat = data["vat"]
    _sheet(wb, "TVA", ["Libellé", "Montant"], [("TVA collectée (4431)", vat["collected"]), ("TVA déductible (4452)", vat["deductible"]), ("TVA À PAYER", vat["due"])], money_cols=(2,)).column_dimensions["A"].width = 32
    cash = data["cash"]
    _sheet(wb, "Trésorerie", ["Compte", "Intitulé", "Ouverture", "Entrées", "Sorties", "Solde"], cash["rows"], money_cols=(3, 4, 5, 6), total_cols=(3, 4, 5, 6))
    wb.save(path)
    return path


def export_accounting_pdf(data: dict, path: str | Path | None = None, journal_limit: int = 200) -> Path:
    config.ensure_directories()
    shop = settings_service.get_shop_info()
    currency = shop.currency or "FCFA"
    if path is None:
        path = config.EXPORT_DIR / f"comptabilite_{datetime.now():%Y%m%d_%H%M%S}.pdf"
    path = Path(path)
    styles = getSampleStyleSheet()
    styles["Heading3"].keepWithNext = 1
    name = str(shop.name or "")
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), topMargin=15 * mm, bottomMargin=18 * mm,
                            leftMargin=15 * mm, rightMargin=15 * mm, title=f"Comptabilité {name}", author=name)
    money = lambda v: format_money(v, currency) if v else ""  # noqa: E731
    story = [Paragraph(f"<b>{name}</b>", styles["Title"]),
             Paragraph(f"Comptabilité du {data['start']:%d/%m/%Y} au {data['end']:%d/%m/%Y}", styles["Heading2"]), Spacer(1, 6)]
    if data["unbalanced"]:
        story.append(Paragraph(f"<font color='#b91c1c'>Attention : écritures déséquilibrées : {', '.join(data['unbalanced'][:10])}</font>", styles["Normal"]))

    if any(code == "311" for code, _l, _v in data["sheet"]["liabilities"]):
        story.append(Paragraph("<font color='#b45309'>Le stock comptable est négatif : enregistrez le stock initial (Comptabilité → Écritures manuelles).</font>", styles["Normal"]))
        story.append(Spacer(1, 6))

    def add(title, header, rows, widths, right_from=1, bold_where=None, limit=None):
        if not rows:
            return
        shown = rows[:limit] if limit else rows
        story.append(Paragraph(f"<b>{title}</b>", styles["Heading3"]))
        t = _table([header] + [list(r) for r in shown], col_widths=widths)
        cmds = [("ALIGN", (right_from, 0), (-1, -1), "RIGHT"), ("FONTSIZE", (0, 0), (-1, -1), 8)]
        if bold_where:
            cmds += [("FONTNAME", (0, i), (-1, i), "Helvetica-Bold") for i, r in enumerate(shown, start=1) if bold_where(r)]
        t.setStyle(TableStyle(cmds))
        story.append(t)
        if limit and len(rows) > limit:
            story.append(Paragraph(f"<font size=8 color='#64748b'>… {len(rows) - limit} ligne(s) supplémentaire(s) dans l'export Excel.</font>", styles["Normal"]))
        story.append(Spacer(1, 12))

    ir = acc.income_rows(data)
    add("Compte de résultat", ["Libellé", "Montant"], [(l, format_money(a, currency)) for l, a, _b in ir], [170 * mm, 60 * mm],
        bold_where=lambda r: r[0].isupper() or r[0].startswith("MARGE"))
    sheet = data["sheet"]
    side = [(a, l, format_money(v, currency)) for a, l, v in sheet["assets"]] + [("", "TOTAL ACTIF", format_money(sheet["total_assets"], currency))]
    add("Bilan — Actif", ["Compte", "Intitulé", "Montant"], side, [25 * mm, 145 * mm, 60 * mm], right_from=2, bold_where=lambda r: r[1].startswith("TOTAL"))
    side = [(a, l, format_money(v, currency)) for a, l, v in sheet["liabilities"]] + [("", "TOTAL PASSIF", format_money(sheet["total_liabilities"], currency))]
    add("Bilan — Passif", ["Compte", "Intitulé", "Montant"], side, [25 * mm, 145 * mm, 60 * mm], right_from=2, bold_where=lambda r: r[1].startswith("TOTAL"))
    vat = data["vat"]
    add("TVA", ["Libellé", "Montant"], [("TVA collectée (4431)", format_money(vat["collected"], currency)), ("TVA déductible (4452)", format_money(vat["deductible"], currency)),
                                       ("TVA À PAYER", format_money(vat["due"], currency))], [170 * mm, 60 * mm], bold_where=lambda r: r[0].startswith("TVA À"))
    add("Trésorerie", ["Compte", "Intitulé", "Ouverture", "Entrées", "Sorties", "Solde"],
        [(r[0], r[1], format_money(r[2], currency), format_money(r[3], currency), format_money(r[4], currency), format_money(r[5], currency)) for r in data["cash"]["rows"]],
        [20 * mm, 70 * mm, 35 * mm, 35 * mm, 35 * mm, 35 * mm], right_from=2)
    story.append(PageBreak())
    add("Balance générale", ["Compte", "Intitulé", "Ouv. D", "Ouv. C", "Mvt D", "Mvt C", "Solde D", "Solde C"],
        [(r[0], r[1]) + tuple(money(v) for v in r[2:]) for r in data["balance"]],
        [18 * mm, 70 * mm, 28 * mm, 28 * mm, 28 * mm, 28 * mm, 28 * mm, 28 * mm], right_from=2)
    story.append(PageBreak())
    add("Journal général", ["Date", "Jnl", "Pièce", "Compte", "Libellé", "Débit", "Crédit"],
        [(r[0], r[1], r[2], f"{r[3]} {r[4]}", r[5], money(r[6]), money(r[7])) for r in data["journal"]],
        [30 * mm, 12 * mm, 34 * mm, 58 * mm, 60 * mm, 32 * mm, 32 * mm], right_from=5, limit=journal_limit)
    footer = lambda c, d: _footer(c, d, name)  # noqa: E731
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path
